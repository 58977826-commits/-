"""落地包 Word / Excel 导出（PDF M6）。"""
from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .matcher import MatchResult, ProjectProfile

try:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt
except ImportError:  # pragma: no cover
    Document = None  # type: ignore[misc, assignment]


_HEADER_FILL = PatternFill("solid", fgColor="FFEDEF")
_HEADER_FONT = Font(bold=True, color="B8000E")


def _autosize_columns(ws, max_width: int = 48) -> None:
    for col_idx, column_cells in enumerate(ws.columns, 1):
        length = 0
        for cell in column_cells:
            val = str(cell.value) if cell.value is not None else ""
            length = max(length, min(len(val), max_width))
        ws.column_dimensions[get_column_letter(col_idx)].width = max(12, min(length + 2, max_width))


def _write_sheet(ws, title: str, headers: list[str], rows: list[list[Any]]) -> None:
    ws.title = title[:31]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for row in rows:
        ws.append(row)
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    _autosize_columns(ws)


def export_landing_pack_excel(
    pack_dir: Path,
    profile: ProjectProfile,
    result: MatchResult,
) -> Path:
    """生成多 sheet Excel 落地包。"""
    wb = Workbook()
    wb.remove(wb.active)

    meta_rows = [
        ["画像ID", profile.画像ID],
        ["客户名称", profile.客户名称],
        ["行业", profile.行业 or "—"],
        ["项目类型", ", ".join(profile.项目类型)],
        ["用户规模", profile.用户规模 or "—"],
        ["参考样板", profile.参考样板项目 or "—"],
        ["推荐服务项数", len(result.service_items)],
        ["推荐组件数", len(result.components)],
        ["识别风险数", len(result.risks)],
    ]
    ws0 = wb.create_sheet("项目信息")
    ws0.append(["字段", "值"])
    for r in meta_rows:
        ws0.append(r)
    _autosize_columns(ws0)

    _write_sheet(
        wb.create_sheet("推荐服务项"),
        "推荐服务项",
        ["服务项ID", "名称", "推荐理由", "规则ID"],
        [[s.item_id, s.name, s.reason, s.rule_id] for s in result.service_items],
    )
    _write_sheet(
        wb.create_sheet("推荐组件"),
        "推荐组件",
        ["组件编号", "组件名称", "来源项目", "推荐理由", "规则ID"],
        [[c.item_id, c.name, c.source_project, c.reason, c.rule_id] for c in result.components],
    )
    _write_sheet(
        wb.create_sheet("RACI"),
        "RACI",
        ["活动", "客户", "菲信运营", "PM", "账务TEM", "运营商"],
        [
            ["项目启动", "A", "R", "C", "I", "I"],
            ["数据导入/TEM", "I", "R", "C", "R", "I"],
            ["月度对账", "I", "R", "C", "R", "C"],
            ["事件闭环", "C", "R", "A", "I", "C"],
            ["月报汇报", "A", "R", "C", "C", "I"],
        ],
    )
    _write_sheet(
        wb.create_sheet("事件字典"),
        "事件字典",
        ["事件类型", "事件动作", "说明", "责任角色", "SLA"],
        [
            ["套餐变更", "升档/降档", "套餐调整", "运营", "5工作日"],
            ["改号", "换号", "号码变更", "运营+运营商", "10工作日"],
            ["离职", "销号/归还", "通信与资产闭环", "运营+HR", "按项目SLA"],
            ["维修", "换机", "终端维修闭环", "运营", "按项目SLA"],
        ],
    )
    _write_sheet(
        wb.create_sheet("风险清单"),
        "风险清单",
        ["风险ID", "风险描述", "关联规则", "确认方", "状态"],
        [[r.item_id, r.name, r.rule_id, "PM/运营", "待确认"] for r in result.risks],
    )
    diff_rows = []
    for d in result.pending_diffs:
        diff_rows.append([
            d.get("规则ID", "—"),
            d.get("规则类别", "—"),
            d.get("标准规则", "—"),
            d.get("本项目规则", "—"),
            d.get("确认方", "—"),
            "是" if d.get("是否可复用") else "否",
            d.get("备注", "—"),
        ])
    _write_sheet(
        wb.create_sheet("差异确认表"),
        "差异确认表",
        ["规则ID", "规则类别", "标准规则", "本项目规则", "确认方", "是否可复用", "备注"],
        diff_rows or [["—", "—", "—", "—", "—", "—", "暂无差异规则"]],
    )

    sop_rows = []
    for c in result.components:
        if c.item_id.startswith("SOP") or "流程" in c.name:
            sop_rows.append([c.item_id, c.name, c.source_project or "—", c.reason])
    _write_sheet(
        wb.create_sheet("SOP索引"),
        "SOP索引",
        ["组件编号", "SOP名称", "来源项目", "推荐原因"],
        sop_rows or [["—", "—", "—", "—"]],
    )

    out = pack_dir / f"{profile.客户名称}_落地包.xlsx"
    wb.save(out)
    return out


def _add_table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = h
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = str(val) if val is not None else ""


def export_landing_pack_word(
    pack_dir: Path,
    profile: ProjectProfile,
    result: MatchResult,
    *,
    brief_excerpt: str = "",
) -> Path:
    """生成 Word 落地包初稿。"""
    if Document is None:
        raise RuntimeError("缺少 python-docx 依赖，请执行 pip install python-docx")

    doc = Document()
    title = doc.add_heading(f"{profile.客户名称} 项目落地包", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    p = doc.add_paragraph()
    p.add_run(f"画像ID：{profile.画像ID}").bold = True
    doc.add_paragraph(f"项目类型：{', '.join(profile.项目类型) or '—'}")
    doc.add_paragraph(f"用户规模：{profile.用户规模 or '—'}")
    doc.add_paragraph(f"参考样板：{profile.参考样板项目 or '—'}")

    doc.add_heading("一、推荐服务项", level=1)
    _add_table(
        doc,
        ["服务项ID", "名称", "推荐理由"],
        [[s.item_id, s.name, s.reason] for s in result.service_items] or [["—", "—", "—"]],
    )

    doc.add_heading("二、推荐组件", level=1)
    _add_table(
        doc,
        ["组件编号", "组件名称", "来源", "推荐理由"],
        [[c.item_id, c.name, c.source_project or "—", c.reason] for c in result.components] or [["—", "—", "—", "—"]],
    )

    doc.add_heading("三、RACI 矩阵", level=1)
    _add_table(
        doc,
        ["活动", "客户", "菲信运营", "PM", "账务/TEM"],
        [
            ["项目启动", "A", "R", "C", "I"],
            ["数据导入", "I", "R", "C", "R"],
            ["月度对账", "I", "R", "C", "R"],
            ["月报汇报", "A", "R", "C", "C"],
        ],
    )

    doc.add_heading("四、风险清单", level=1)
    _add_table(
        doc,
        ["风险描述", "关联规则", "状态"],
        [[r.name, r.rule_id, "待确认"] for r in result.risks] or [["—", "—", "—"]],
    )

    doc.add_heading("五、项目差异确认", level=1)
    diff_rows = [
        [d.get("规则ID", "—"), d.get("标准规则", "—"), d.get("本项目规则", "—"), d.get("确认方", "—")]
        for d in result.pending_diffs
    ]
    _add_table(doc, ["规则ID", "标准规则", "本项目规则", "确认方"], diff_rows or [["—", "—", "—", "—"]])

    if brief_excerpt:
        doc.add_heading("六、AI Brief 摘要", level=1)
        for line in brief_excerpt.splitlines()[:40]:
            para = doc.add_paragraph(line)
            para.paragraph_format.space_after = Pt(2)

    doc.add_paragraph("")
    doc.add_paragraph("（本文件为系统生成的项目落地包初稿，请 PM/运营确认后对外使用。）")

    out = pack_dir / f"{profile.客户名称}_落地包.docx"
    doc.save(out)
    return out


def export_landing_pack_office(
    pack_dir: Path,
    profile: ProjectProfile,
    result: MatchResult,
    *,
    brief_excerpt: str = "",
) -> dict[str, str]:
    """同时导出 Excel 与 Word，返回文件路径。"""
    pack_dir = Path(pack_dir)
    xlsx = export_landing_pack_excel(pack_dir, profile, result)
    paths = {"xlsx": str(xlsx)}
    try:
        docx = export_landing_pack_word(pack_dir, profile, result, brief_excerpt=brief_excerpt)
        paths["docx"] = str(docx)
    except RuntimeError:
        paths["docx"] = ""
    return paths
